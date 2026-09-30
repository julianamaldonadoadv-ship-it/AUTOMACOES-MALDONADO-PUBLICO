# -*- coding: utf-8 -*-
"""
Trava de custas antes do protocolo (POP-CJ-PROT-001 x POP-CJ-006)
==================================================================

Nasceu de um caso real: a declaratoria da Sra. Cliente G (0000021-76.2026.8.22.0001)
foi protocolada em 22/09/2026 as 10:49 SEM o recolhimento das custas. As tarefas
de custas nasceram depois do protocolo — RECOLHER CUSTAS as 10:53 e SOLICITAR
DOCUMENTACAO ("documentos para emenda: comprovante de residencia e juntada de
custas") as 10:57. O escritorio gastou emenda a inicial, prazo e retrabalho num
item que e' checklist.

A pergunta que a trava responde, antes de a peca ir para o protocolo:

    1. E' peticao INICIAL?                      -> se nao for, a trava nao se aplica
    2. Tem topico/pedido de GRATUIDADE?
       SIM -> os documentos que sustentam a gratuidade estao na pasta?
              (declaracao de hipossuficiencia e a prova de renda; art. 99, § 2o,
              permite ao juizo indeferir quando ha elementos nos autos que
              contrariam a alegacao — em produtor rural isso e' regra, nao excecao)
       NAO -> a GUIA foi emitida pela Controladoria (tarefa RECOLHER CUSTAS)?
              o COMPROVANTE de pagamento esta juntado?
              a CONTROLADORIA foi acionada, e o CS levou o boleto ao cliente?
    3. Os documentos obrigatorios da tese estao na pasta? (o mesmo rol da skill
       `inicial-anexos` — no caso da Cliente G faltava tambem comprovante de residencia)

Duas camadas, de proposito:

    CAMADA 1 (preventiva) — `custas conferir peca.docx`
        roda ANTES do protocolo, le a peca, cruza com a pasta do cliente na ZEUS
        e com o ADVBOX. Status BLOQUEADO sai com codigo de saida 1, para poder
        travar script/rotina.

    CAMADA 2 (rede de seguranca) — `custas auditar --dias 7`
        varre as INICIAIS que ja' foram protocoladas e acusa aquela que foi sem
        custas e sem gratuidade. Pega o que escapou da camada 1 — inclusive peca
        redigida fora da automacao, que e' a maioria.

Calibrado contra a peca real da Sra. Cliente G (nao desfazer):

  - **A peca ANUNCIAVA a guia e a guia nao existia.** O rol de documentos da
    inicial termina com "... e guia de recolhimento das custas iniciais". Ler
    isso como prova de recolhimento liberaria justamente o protocolo que deu
    errado. Aqui a mencao conta ao contrario: e' uma promessa que a trava exige
    conferir no arquivo (pendencia GUIA_ANUNCIADA_NAO_JUNTADA).
  - A peca nao tinha topico de gratuidade: caiu na via das custas, como devia.
  - Marcador de pendencia que sobrou na peca ([PENDENTE], [[INSERIR ...]])
    bloqueia o protocolo — e' o mesmo guard-rail da skill `timbrado`, que proibe
    "conforme documento anexo" sem documento.

Guard-rails (os mesmos do resto do repositorio):
  - Somente leitura. Nenhuma tarefa e' criada sem `--criar-tarefa` e sem "s/N".
  - A automacao NUNCA protocola e nunca decide sozinha que a gratuidade cabe:
    ela confere se o pedido esta na peca e se a prova esta na pasta.
  - Leitura de texto e' INDICIO. Peca que a trava nao consegue classificar sobe
    como A CONFERIR (nao como liberada).
  - Ha' saida de excecao (`--apesar-de "motivo"`): prazo fatal no mesmo dia
    acontece. O motivo entra por escrito na tarefa, com data e autor — o que a
    trava impede e' o protocolo sem custas SILENCIOSO.
"""
import os
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'INTEGRACOES'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'config'))

try:
    import advbox_integration as advbox
except Exception:          # autoteste roda sem credencial
    advbox = None
try:
    import roteamento_controller as roteamento
except Exception:
    roteamento = None
try:
    import prazos
except Exception:
    prazos = None
try:
    import equipe
except ImportError:
    equipe = None


def _norm(texto):
    t = unicodedata.normalize('NFKD', texto or '').encode('ascii', 'ignore').decode().lower()
    t = re.sub(r'<[^>]+>', ' ', t)
    return re.sub(r'\s+', ' ', t)


# ============================================================
# 1. LEITURA DA PECA
# ============================================================

def ler_peca(caminho):
    """Texto da peca. Aceita .docx (corpo + tabelas), .pdf, .txt e .md."""
    ext = os.path.splitext(caminho)[1].lower()
    if ext == '.docx':
        from docx import Document
        doc = Document(caminho)
        partes = [p.text for p in doc.paragraphs]
        for tabela in doc.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    partes.append(celula.text)
        return '\n'.join(partes)
    if ext == '.pdf':
        import fitz
        with fitz.open(caminho) as pdf:
            return '\n'.join(pagina.get_text() for pagina in pdf)
    with open(caminho, encoding='utf-8', errors='ignore') as fh:
        return fh.read()


# ============================================================
# 2. E' INICIAL?
# ============================================================
# Score, nao chave unica: nenhuma expressao isolada distingue inicial de
# peticao intermediaria (ate "termos em que pede deferimento" fecha as duas).

_INDICIOS_INICIAL = (
    (r'\bpeticao inicial\b', 2, 'menciona "peticao inicial"'),
    (r'propor a presente|ajuizar a presente|propositura da presente|'
     r'vem .{0,80}propor|vem .{0,80}ajuizar', 3, 'verbo de propositura da acao'),
    (r'requer(?:er)? a citacao|citacao (?:da|do) (?:parte )?(?:requerid|r[eé]|execut)',
     3, 'pedido de citacao'),
    (r'da-se a causa o valor|valor da causa', 2, 'atribui valor a causa'),
    (r'protesta provar o alegado|provar o alegado por todos os meios',
     1, 'protesto generico de provas'),
    (r'distribuicao por dependencia', 2, 'pede distribuicao por dependencia'),
    (r'excelentissim[oa].{0,120}(?:juiz|juiza|vara|juizado)', 1,
     'enderecamento a juizo de 1o grau'),
    (r'qualificad[oa].{0,40}(?:residente|domiciliad|inscrit[oa] no cpf|portador)',
     1, 'qualificacao completa da parte'),
)

# O que denuncia peca que NAO e' inicial. Peso alto: errar aqui faz a trava
# bloquear manifestacao e virar ruido — e trava que vira ruido e' desligada.
_INDICIOS_NAO_INICIAL = (
    (r'nos autos (?:do processo )?em epigrafe|nos autos em epigrafe|'
     r'ja (?:devidamente )?qualificad[oa] nos autos|'
     r'ja qualificad[oa]|nos autos do processo em referencia', 4,
     'parte ja qualificada nos autos'),
    (r'razoes (?:de|do) (?:apelacao|recurso)|contrarrazoes|'
     r'agravo de instrumento|agravo interno|embargos de declaracao|'
     r'recurso especial|recurso extraordinario|recurso inominado', 4,
     'peca recursal'),
    (r'embargos a execucao|excecao de pre-?executividade|'
     r'impugnacao ao cumprimento|impugnacao a penhora', 3, 'defesa em execucao'),
    (r'replica|impugnacao a contestacao|alegacoes finais|memoriais|'
     r'especificacao de provas|manifestacao sobre', 3, 'peca de fase posterior'),
    (r'emenda (?:a|da) inicial|aditamento (?:a|da) inicial', 4, 'emenda/aditamento'),
    (r'egregi[oa] (?:tribunal|camara|turma)|colend[ao]', 2, 'enderecamento a tribunal'),
)


# ============================================================
# 2b. RECURSO — preparo, e nao custa inicial
# ============================================================
# Dispositivos conferidos no texto compilado do CPC que esta em
# BASE_CONHECIMENTO/06 - LEGISLACAO/CPC/ (nao citar de memoria):
#   art. 1.007, caput — preparo NO ATO da interposicao, sob pena de desercao
#   art. 1.007, § 3o  — porte de remessa/retorno dispensado em autos eletronicos
#   art. 1.007, § 4o  — nao comprovado, intima para recolher EM DOBRO
#   art. 1.007, § 5o  — vedada a complementacao no recolhimento em dobro
#   art. 1.017, § 1o  — o agravo vai acompanhado do comprovante do pagamento
#   art. 1.017, § 5o  — autos eletronicos dispensam as copias do inciso I
#   art. 1.023        — embargos de declaracao NAO se sujeitam a preparo
#   art. 98, § 1o, I  — a gratuidade compreende as taxas e custas judiciais
#
# A diferenca pratica em relacao a inicial: la o erro vira emenda (art. 321,
# 15 dias, e so entao indeferimento); aqui vira DESERCAO, e o § 5o proibe
# complementar o que foi recolhido em dobro. Recurso perdido por preparo nao
# tem segunda chance.

RECURSOS = {
    'agravo-instrumento': (r'agravo de instrumento', True),
    'apelacao': (r'\bapelacao\b|razoes de apelacao|recurso de apelacao', True),
    'inominado': (r'recurso inominado', True),
    # "REsp 1.061.530/RS" e' CITACAO do Tema 28, nao a peca: por isso so
    # "recurso especial" por extenso conta (a inicial da Cliente G virava RESP).
    'especial': (r'recurso especial', True),
    'extraordinario': (r'recurso extraordinario', True),
    # Sem preparo, por dispositivo ou por natureza:
    'embargos-declaracao': (r'embargos de declaracao', False),
    'agravo-interno': (r'agravo interno|agravo regimental', False),
    'contrarrazoes': (r'(?:apresentar|oferecer|ofertar|apresenta|oferece)\s+'
                      r'(?:as\s+)?contrarrazoes|contrarrazoes (?:ao|aos|de|da|do)\b', False),
}
# Recurso cujo preparo depende do regimento de custas do tribunal — a trava nao
# afirma nem nega, manda conferir.
_COM_PREPARO_SEMPRE = ('agravo-instrumento', 'apelacao', 'inominado',
                       'especial', 'extraordinario')


# Defesa em execucao nao e' recurso nem inicial comum: se ha custas depende do
# regimento do tribunal. A trava nao afirma nem nega — manda conferir.
_DEFESA_EXECUCAO = r'embargos a execucao|excecao de pre-?executividade|impugnacao ao cumprimento'

# So o CABECALHO decide a natureza da peca. No corpo, toda peca cita agravo e
# apelacao (historico do processo, precedente): procurar no texto inteiro fazia
# a contrarrazoes virar agravo de instrumento e os embargos a execucao virarem
# apelacao — conferido nas pecas reais em 22/09/2026.
_JANELA_CABECALHO = 2500


# O que identifica a peca e' o VERBO DE APRESENTACAO dela ("vem ... interpor o
# presente recurso especial"), nao o nome do recurso solto. Sem isso:
#   - a contrarrazoes do Cliente R virava agravo, porque o enderecamento e'
#     "EXCELENTISSIMO RELATOR DO AGRAVO DE INSTRUMENTO 0000029-00...";
#   - a peticao de tutela incidental da Sra. Cliente K virava agravo, porque um
#     dos topicos e' "preservacao do direito de interpor agravo de instrumento".
# O intervalo entre "vem" e o verbo NAO pode excluir ponto: no cabecalho do REsp
# da Sra. Cliente K vem "com fundamento no art. 105, III, ... interpor o presente
# recurso especial" — o "art." cortava o casamento.
_VERBO_APRESENTACAO = (r'(?:vem|venho|vimos)\b.{0,240}?'
                       r'(?:interpor|interpoe|opor|opoe|apresentar|apresenta|'
                       r'oferecer|oferece|ofertar)\s+'
                       r'(?:o\s+|a\s+|os\s+|as\s+)?(?:presente[s]?\s+)?')


def identificar_recurso(texto_normalizado):
    """Devolve (slug, tem_preparo, confianca) do recurso, ou (None, False, None).

    confianca: 'alta'  — a peca diz que esta interpondo/opondo/apresentando
               'baixa' — o nome do recurso apareceu no cabecalho, sem verbo.
    Confianca baixa nao bloqueia protocolo: sobe como conferencia.
    """
    cabecalho = texto_normalizado[:_JANELA_CABECALHO]

    # 1a passada: o verbo de apresentacao manda.
    declarados = []
    for slug, (padrao, com_preparo) in RECURSOS.items():
        m = re.search(_VERBO_APRESENTACAO + r'(?:' + padrao + r')', cabecalho)
        if m:
            declarados.append((m.start(), slug, com_preparo))
    if declarados:
        declarados.sort()
        return declarados[0][1], declarados[0][2], 'alta'

    # 2a passada: so o nome, sem verbo — indicio, nao conclusao.
    achados = []
    for slug, (padrao, com_preparo) in RECURSOS.items():
        m = re.search(padrao, cabecalho)
        if m:
            achados.append((m.start(), slug, com_preparo))
    if not achados:
        return None, False, None
    for preferido in ('contrarrazoes', 'embargos-declaracao', 'agravo-interno'):
        for _, slug, com_preparo in achados:
            if slug == preferido:
                return slug, com_preparo, 'baixa'
    achados.sort()
    return achados[0][1], achados[0][2], 'baixa'


def eh_defesa_execucao(texto_normalizado):
    return bool(re.search(_DEFESA_EXECUCAO, texto_normalizado[:_JANELA_CABECALHO]))


def classificar_peca(texto):
    """Diz se a peca e' inicial. Devolve tipo, indicios e contra-indicios."""
    t = _norm(texto)
    pro, contra = [], []
    peso_pro = peso_contra = 0
    for padrao, peso, rotulo in _INDICIOS_INICIAL:
        if re.search(padrao, t):
            pro.append(rotulo)
            peso_pro += peso
    for padrao, peso, rotulo in _INDICIOS_NAO_INICIAL:
        if re.search(padrao, t):
            contra.append(rotulo)
            peso_contra += peso

    recurso, com_preparo, confianca_recurso = identificar_recurso(t)

    if eh_defesa_execucao(t) and not recurso:
        return {'tipo': 'defesa-execucao', 'recurso': None, 'preparo': None,
                'indicios': pro, 'contra_indicios': contra,
                'peso': peso_pro, 'peso_contra': peso_contra}
    if recurso and confianca_recurso == 'alta':
        # Recurso e' peca autonoma para a trava: nao paga custa inicial, paga
        # preparo. Antes caia em 'outra' e saia sem conferencia nenhuma.
        return {'tipo': 'recurso', 'recurso': recurso, 'preparo': com_preparo,
                'confianca_recurso': confianca_recurso,
                'indicios': pro, 'contra_indicios': contra,
                'peso': peso_pro, 'peso_contra': peso_contra}
    if peso_contra >= 4 and peso_contra >= peso_pro:
        tipo = 'outra'
    elif peso_pro >= 5 and peso_pro > peso_contra:
        tipo = 'inicial'
    elif peso_pro >= 3 and peso_contra == 0:
        tipo = 'inicial'
    else:
        tipo = 'indefinido'
    # Nome de recurso citado numa inicial (o "REsp 1.061.530/RS" do Tema 28) nao
    # faz da inicial um recurso: fora do tipo 'recurso', so o achado de confianca
    # alta viaja adiante.
    citado = recurso if confianca_recurso == 'baixa' else None
    if confianca_recurso != 'alta':
        recurso, com_preparo = None, None
    return {'tipo': tipo, 'recurso': recurso, 'preparo': com_preparo,
            'recurso_citado': citado, 'confianca_recurso': confianca_recurso,
            'indicios': pro, 'contra_indicios': contra,
            'peso': peso_pro, 'peso_contra': peso_contra}


# ============================================================
# 3. O QUE A PECA DIZ SOBRE AS CUSTAS
# ============================================================

_GRATUIDADE = (
    r'gratuidade (?:da justica|judiciaria|de justica)',
    r'justica gratuita',
    r'assistencia judiciaria gratuita',
    r'beneficio da (?:gratuidade|justica gratuita|assistencia judiciaria)',
    r'hipossuficien(?:te|cia)',
    r'lei (?:n\.?o?\s*)?1\.?060',
    r'isencao (?:das|de) custas',
)
# art. 98 sozinho nao conta: o numero aparece em citacao de outro contexto.
_GRATUIDADE_ART = r'art(?:igo)?\.?\s*98\b'

_DIFERIMENTO = (
    r'diferimento (?:das|do recolhimento das) custas',
    r'diferir o recolhimento',
    r'recolhimento (?:das custas )?ao final',
    r'custas ao final',
)
_PARCELAMENTO = (
    r'parcelamento das custas',
    r'parcelar as custas',
    r'recolhimento parcelado',
)
_RECOLHIMENTO = (
    r'custas (?:iniciais )?(?:ja )?recolhidas',
    r'comprovante de (?:recolhimento|pagamento) das custas',
    r'guia de custas',
    r'recolhimento das custas iniciais',
    r'custas devidamente recolhidas',
)
# Marcador que o proprio escritorio deixa na minuta quando falta documento
# (skill `timbrado` e `inicial-anexos`). Peca que vai para protocolo com
# marcador vivo ainda esta em producao — nao esta pronta.
_MARCADORES_PENDENTES = (
    r'\[pendente\]',
    r'\[\[\s*inserir[^\]]{0,80}\]\]',
    r'\[\[\s*texto do print[^\]]{0,40}\]\]',
    r'inserir uma imagem\.',
    r'imagem do requerimento',
    r'inserir grafico',
    r'\bxxx+\b',
)

_VALOR_CAUSA_EXPRESSO = re.compile(
    r'(?:da-se (?:a|à) causa o valor de|valor da causa\s*:?)\D{0,30}'
    r'r\$\s*([\d\.]{1,15},\d{2})')
# Reserva. Deixado por ultimo de proposito: "valor de R$ ..." solto pega o valor
# do CONTRATO (na peca da Cliente G, R$ 787.500,00 contra a causa de R$ 1.000,00).
_VALOR_CAUSA = re.compile(
    r'(?:valor (?:da causa|de))\D{0,30}r\$\s*([\d\.]{1,15},\d{2})')


def _trechos(texto_original, t_norm, padroes):
    achados = []
    for padrao in padroes:
        for m in re.finditer(padrao, t_norm):
            ini = max(0, m.start() - 70)
            achados.append(t_norm[ini:m.end() + 90].strip())
            break
    return achados


def analisar_peca(texto):
    """Le a peca e devolve o que ela diz (ou nao diz) sobre as custas."""
    t = _norm(texto)
    classe = classificar_peca(texto)

    gratuidade = any(re.search(p, t) for p in _GRATUIDADE)
    # "art. 98" so conta como gratuidade se vier perto de custas/beneficio
    if not gratuidade and re.search(_GRATUIDADE_ART, t):
        for m in re.finditer(_GRATUIDADE_ART, t):
            redor = t[max(0, m.start() - 120):m.end() + 120]
            if 'custas' in redor or 'gratuid' in redor or 'beneficio' in redor:
                gratuidade = True
                break

    valor = _VALOR_CAUSA_EXPRESSO.search(t) or _VALOR_CAUSA.search(t)
    return {
        'tipo': classe['tipo'],
        'recurso': classe.get('recurso'),
        'preparo': classe.get('preparo'),
        'confianca_recurso': classe.get('confianca_recurso'),
        'classificacao': classe,
        'gratuidade': gratuidade,
        'gratuidade_trechos': _trechos(texto, t, _GRATUIDADE),
        'diferimento': any(re.search(p, t) for p in _DIFERIMENTO),
        'parcelamento': any(re.search(p, t) for p in _PARCELAMENTO),
        # A peca DIZER que a guia vai anexa nao e' prova de que foi recolhida —
        # e' promessa a conferir. Foi assim que o 0000021-76 passou.
        'anuncia_guia': any(re.search(p, t) for p in _RECOLHIMENTO),
        'marcadores_pendentes': [m.group(0)[:60] for p in _MARCADORES_PENDENTES
                                 for m in re.finditer(p, t)][:12],
        'valor_causa': valor.group(1) if valor else None,
    }


# ============================================================
# 4. DOCUMENTOS QUE SUSTENTAM CADA VIA
# ============================================================
# Classificacao por nome de arquivo — palpite, nao laudo (mesma regra do
# `anexos_inicial`). Arquivo com nome ruim nao casa e sobe como faltante: a
# controller confere na pasta antes de concluir.

DOCS_GRATUIDADE = (
    ('declaracao-hipossuficiencia', 'Declaracao de hipossuficiencia', True,
     ('declaracao de hipossuficiencia', 'hipossuficiencia', 'declaracao de pobreza',
      'decl hipossuficiencia', 'declaracao hipossuficiencia')),
    ('prova-renda', 'Prova da renda (IRPF, CNIS, extrato, contracheque)', False,
     ('irpf', 'imposto de renda', 'declaracao de imposto', 'cnis', 'extrato bancario',
      'contracheque', 'holerite', 'comprovante de renda', 'extrato inss',
      'cadunico', 'cad unico')),
)

DOCS_CUSTAS = (
    ('guia-custas', 'Guia/boleto das custas iniciais', True,
     ('guia', 'boleto', 'dare', 'daje', 'custas iniciais', 'custas')),
    ('comprovante-custas', 'Comprovante de pagamento das custas', True,
     ('comprovante de pagamento', 'comprovante custas', 'comprovante de recolhimento',
      'pagamento custas', 'recibo', 'quitacao', 'comprovante')),
)


def conferir_documentos(nomes, rol):
    """Casa os nomes de arquivo da pasta contra um rol. Devolve achados e faltas."""
    normalizados = [(n, _norm(n)) for n in nomes or ()]
    ok, faltando = [], []
    for slug, rotulo, obrigatorio, termos in rol:
        casados = [n for n, alvo in normalizados if any(term in alvo for term in termos)]
        if casados:
            ok.append({'slug': slug, 'rotulo': rotulo, 'arquivos': casados})
        else:
            faltando.append({'slug': slug, 'rotulo': rotulo, 'obrigatorio': obrigatorio})
    return {'ok': ok, 'faltando': faltando}


# ============================================================
# 5. O QUE O ADVBOX JA REGISTRA
# ============================================================

_EV_GUIA = (r'recolher custas', r'emitir (?:as )?custas', r'guia de custas',
            r'emissao (?:da|de) guia', r'boleto (?:das|de) custas')
_EV_COMPROVANTE = (r'comprovante de pagamento', r'comprovante de recolhimento',
                   r'custas (?:pagas|recolhidas|quitadas)', r'pagamento (?:das )?custas confirmado')
_EV_CS = (r'contato com cliente', r'enviar boleto', r'boleto referente',
          r'avisar cliente')
_EV_EMENDA = (r'emenda a inicial', r'emenda da inicial', r'documentos para emenda')
_EV_ERRO = (r'erro no protocolo',)
_EV_PROTOCOLO = (r'protocolo d-?3', r'protocolo d-?2', r'protocolo d-?1',
                 r'distribuicao\b', r'peca aprovada para protocolo',
                 r'protocolo - prazo fatal')


def _casa_evento(item, padroes):
    alvo = _norm((item.get('task') or '') + ' ' + (item.get('comments') or ''))
    return any(re.search(p, alvo) for p in padroes)


def situacao_advbox(lawsuit_id, historico=None, pendentes=None):
    """
    Le o historico do processo e diz o que ja' aconteceu de custas.

    ATENCAO: `GET /history` devolve ~20 itens e NAO pagina (CLAUDE.md) — em
    processo novo cobre tudo, em processo antigo e' piso, nao total. Por isso
    "nao encontrei" nunca vira "nao existe": vira pendencia a conferir.
    """
    if historico is None:
        historico = advbox.listar_historico(lawsuit_id) or []
    if pendentes is None:
        try:
            pendentes = advbox.listar_tarefas(lawsuit_id=lawsuit_id) or []
        except Exception:
            pendentes = []

    def _primeiro(padroes):
        for item in historico:
            if _casa_evento(item, padroes):
                return item
        return None

    # /posts devolve o que esta em aberto; o que esta no /history e nao esta la'
    # foi concluido. E' inferencia, e sai rotulada como tal no relatorio.
    tarefas_abertas = {_norm(t.get('task') or '') for t in pendentes}

    guia = _primeiro(_EV_GUIA)
    return {
        'guia': guia,
        'guia_pendente': bool(guia) and any('recolher custas' in a for a in tarefas_abertas),
        'comprovante': _primeiro(_EV_COMPROVANTE),
        'cs_acionado': _primeiro(_EV_CS),
        'emenda': _primeiro(_EV_EMENDA),
        'erro_protocolo': _primeiro(_EV_ERRO),
        'protocolo': _primeiro(_EV_PROTOCOLO),
        'eventos': historico,
        'pendentes': pendentes,
        'cobertura_parcial': len(historico) >= 20,
    }


def documentos_advbox(cliente_id):
    """Anexos do cliente no ADVBOX que casam com guia/comprovante de custas."""
    if not cliente_id or not advbox:
        return []
    try:
        docs = advbox.listar_documentos(cliente_id=cliente_id) or []
    except Exception:
        return []
    achados = []
    for d in docs:
        nome = _norm(d.get('name') or d.get('file_name') or '')
        if 'custas' in nome or 'guia' in nome or 'darj' in nome or 'boleto' in nome:
            achados.append(d)
    return achados


# ============================================================
# 6. VEREDITO
# ============================================================

STATUS_LIBERADO = 'LIBERADO'
STATUS_BLOQUEADO = 'BLOQUEADO'
STATUS_CONFERIR = 'A CONFERIR'
STATUS_NAO_SE_APLICA = 'NAO SE APLICA'


def _avaliar_recurso(peca, sit, motivo_excecao=None):
    """Recurso paga PREPARO, nao custa inicial — e o erro aqui nao vira emenda.

    art. 1.007, caput: o preparo se comprova NO ATO da interposicao, sob pena de
    DESERCAO; § 4o: nao comprovado, o recorrente e intimado para recolher EM
    DOBRO; § 5o: e' vedada a complementacao nesse recolhimento em dobro. Ou seja:
    na inicial o escritorio tem 15 dias de emenda (art. 321); no recurso, nao.
    """
    pendencias, observacoes = [], []
    rotulo = (peca.get('recurso') or 'recurso').replace('-', ' ')

    if not peca.get('preparo'):
        observacoes.append(
            f'{rotulo.capitalize()}: sem preparo. '
            'Embargos de declaracao nao se sujeitam a preparo (art. 1.023); agravo '
            'interno corre dentro do recurso ja preparado; contrarrazoes nao e recurso.')
        return {'status': STATUS_NAO_SE_APLICA, 'via': 'preparo',
                'pendencias': pendencias, 'observacoes': observacoes}

    if peca.get('gratuidade'):
        observacoes.append(
            'A peca invoca a gratuidade. Gratuidade DEFERIDA nos autos dispensa o '
            'preparo (art. 98, § 1o, I); gratuidade apenas PEDIDA no recurso nao '
            'dispensa por si — conferir a decisao que a concedeu antes de protocolar.')
        pendencias.append(_pendencia(
            'PREPARO_GRATUIDADE_A_CONFERIR',
            'Confirmar nos autos a decisao que deferiu a gratuidade',
            'Sem decisao de deferimento, o relator trata como preparo nao comprovado '
            'e aplica o art. 1.007, § 4o (recolhimento em dobro).',
            'Controladoria + advogado responsavel', bloqueia=False))
        status = STATUS_CONFERIR
    else:
        tem_guia = bool(sit.get('guia'))
        tem_comprovante = bool(sit.get('comprovante'))
        if not tem_guia:
            pendencias.append(_pendencia(
                'PREPARO_SEM_GUIA',
                f'Guia do preparo do {rotulo} nao foi emitida',
                'art. 1.007, caput: o preparo se comprova no ato da interposicao, '
                'sob pena de desercao.',
                'Controladoria', tarefa='RECOLHER CUSTAS'))
        if not tem_comprovante:
            pendencias.append(_pendencia(
                'PREPARO_SEM_COMPROVANTE',
                f'Comprovante do preparo do {rotulo} nao esta juntado',
                'art. 1.007, § 4o: nao comprovado no ato, intima-se para recolher EM '
                'DOBRO; § 5o veda a complementacao. Recurso perdido por preparo nao '
                'tem emenda — diferente da inicial, que tem 15 dias (art. 321).',
                'Controladoria + CS (cliente paga)', tarefa='RECOLHER CUSTAS'))
        if peca.get('recurso') == 'agravo-instrumento' and not tem_comprovante:
            pendencias.append(_pendencia(
                'AGRAVO_SEM_COMPROVANTE_NA_PECA',
                'O agravo tem de ir instruido com o comprovante do pagamento',
                'art. 1.017, § 1o: acompanha a peticao o comprovante do pagamento das '
                'custas e do porte de retorno, quando devidos. (As copias do inciso I '
                'sao dispensadas em autos eletronicos, § 5o — o comprovante nao.)',
                'Controladoria', tarefa='RECOLHER CUSTAS'))
        if not sit.get('cs_acionado') and not tem_comprovante:
            pendencias.append(_pendencia(
                'CS_NAO_ACIONADO',
                'Cliente nao foi avisado do boleto do preparo',
                'POP-CJ-006: a Controladoria emite e o CS leva ao cliente. No recurso '
                'o prazo e o da interposicao — nao ha janela de emenda.',
                'CS (Anna Lydia / Karla)', tarefa='CONTATO COM CLIENTE'))
        status = STATUS_BLOQUEADO if pendencias else STATUS_LIBERADO

    observacoes.append('Porte de remessa e retorno e dispensado em autos eletronicos '
                       '(art. 1.007, § 3o) — no PJe cobra-se so o preparo.')

    for marcador in peca.get('marcadores_pendentes') or ():
        pendencias.append(_pendencia(
            'MARCADOR_PENDENTE',
            f'A peca ainda tem marcador de pendencia: "{marcador}"',
            'Marcador vivo quer dizer trecho ou documento que ficou faltando.',
            'Advogado responsavel'))
        status = STATUS_BLOQUEADO

    if motivo_excecao and any(p['bloqueia'] for p in pendencias):
        status = STATUS_LIBERADO
        observacoes.append('EXCECAO REGISTRADA: ' + motivo_excecao)
        observacoes.append('Em recurso a excecao e mais grave que na inicial: o § 5o do '
                           'art. 1.007 veda complementar o recolhimento feito em dobro.')

    return {'status': status, 'via': 'preparo', 'pendencias': pendencias,
            'observacoes': observacoes, 'excecao': motivo_excecao}


def _pendencia(codigo, o_que, por_que, quem, tarefa=None, bloqueia=True):
    return {'codigo': codigo, 'o_que': o_que, 'por_que': por_que, 'quem': quem,
            'tarefa': tarefa, 'bloqueia': bloqueia}


def pasta_confere(nome_pasta, nome_cliente):
    """A pasta que o Drive devolveu e' mesmo a pasta de documentos do cliente?

    `google_integration.pasta_do_cliente()` cai numa busca global quando o cliente
    nao esta sob a letra esperada — e ai ela devolve, por exemplo, a pasta de
    PECAS AUTOMACAO ("FULANO - SEM PROCESSO"), que nao tem documento nenhum.
    Bloquear a peca por "falta documento" lendo a pasta errada seria pior que nao
    conferir: a trava perde credibilidade na primeira semana.
    """
    if not nome_pasta:
        return False
    a, b = _norm(nome_pasta), _norm(nome_cliente)
    if a == b:
        return True
    # sufixo de pasta de peca: "- SEM PROCESSO" ou "- 7001234-56.2026.8.22.0001"
    if re.search(r'-\s*(sem processo|\d{7}-\d{2}\.\d{4})', a):
        return False
    return b in a and len(a) - len(b) <= 3


def avaliar(peca=None, arquivos_pasta=None, advbox_situacao=None,
            rol_tese=None, motivo_excecao=None, pasta_confiavel=True):
    """
    Junta as tres leituras (peca, pasta, ADVBOX) e diz se pode protocolar.

    peca              — saida de analisar_peca() (None = peca nao lida)
    arquivos_pasta    — nomes de arquivo da pasta do cliente / do pacote
    advbox_situacao   — saida de situacao_advbox()
    rol_tese          — saida de anexos_inicial.conferir_rol() (documentos da tese)
    motivo_excecao    — texto: libera assim mesmo, registrando o motivo
    """
    pendencias, observacoes = [], []

    if peca is None:
        return {'status': STATUS_CONFERIR, 'via': None, 'pendencias': [],
                'observacoes': ['Peca nao foi lida — a trava conferiu so o ADVBOX.']}

    if peca['tipo'] == 'outra':
        return {'status': STATUS_NAO_SE_APLICA, 'via': None, 'pendencias': [],
                'observacoes': ['Nao e peticao inicial: ' +
                                '; '.join(peca['classificacao']['contra_indicios'][:2])]}

    if peca['tipo'] == 'defesa-execucao':
        return {'status': STATUS_CONFERIR, 'via': 'custas-incidentais', 'pendencias': [
            _pendencia('DEFESA_EXECUCAO_CUSTAS',
                       'Conferir no regimento de custas do tribunal se a defesa em '
                       'execucao tem custas proprias',
                       'Embargos a execucao e excecao de pre-executividade nao entram '
                       'na regra da inicial nem na do preparo — depende do regimento '
                       'de cada tribunal, e a trava nao arbitra isso.',
                       'Controladoria', bloqueia=False)],
            'observacoes': ['Peca de defesa em execucao.']}

    if peca['tipo'] == 'recurso':
        return _avaliar_recurso(peca, sit=advbox_situacao or {},
                                motivo_excecao=motivo_excecao)

    indefinida = peca['tipo'] == 'indefinido'
    if peca['classificacao'].get('recurso_citado'):
        observacoes.append(
            'O cabecalho cita "' + peca['classificacao']['recurso_citado'].replace('-', ' ')
            + '" sem dizer que esta interpondo. Tratei como nao-recurso; se for '
              'recurso, o que vale e o PREPARO (art. 1.007), nao a custa inicial.')
    if indefinida:
        observacoes.append(
            'Nao consegui afirmar que e peticao inicial (indicios fracos). '
            'A conferencia segue como se fosse — confirme antes de protocolar.')

    nomes = list(arquivos_pasta or ())
    sit = advbox_situacao or {}

    # ---- via 1: gratuidade
    if peca['gratuidade']:
        via = 'gratuidade'
        conf = conferir_documentos(nomes, DOCS_GRATUIDADE)
        if arquivos_pasta is None:
            observacoes.append('Pasta do cliente nao foi conferida (rode com --cliente): '
                               'a prova da hipossuficiencia nao foi verificada.')
        else:
            for falta in conf['faltando']:
                if falta['obrigatorio']:
                    pendencias.append(_pendencia(
                        'GRATUIDADE_SEM_DECLARACAO',
                        f"Falta {falta['rotulo']} na pasta do cliente",
                        'Pedido de gratuidade sem a declaracao anexa e indeferido de '
                        'plano ou vira emenda — e a emenda custa o prazo.',
                        'Setor de documentos (INICIAL - ORGANIZAR DOCUMENTOS)'))
                else:
                    pendencias.append(_pendencia(
                        'GRATUIDADE_SEM_PROVA_RENDA',
                        f"Sem {falta['rotulo']}",
                        'Art. 99, § 2o: havendo nos autos elementos que contrariem a '
                        'alegacao, o juizo indefere. Produtor rural com area e '
                        'maquinario cai nessa hipotese com frequencia.',
                        'Advogado responsavel', bloqueia=False))

    # ---- via 2: diferimento / parcelamento
    elif peca['diferimento'] or peca['parcelamento']:
        via = 'diferimento' if peca['diferimento'] else 'parcelamento'
        observacoes.append(
            'A peca pede ' + via + ' das custas. Confirme que o pedido esta em '
            'topico proprio, fundamentado e com prova da situacao financeira — '
            'pedido solto no meio da peca costuma passar batido e virar emenda.')
        pendencias.append(_pendencia(
            'DIFERIMENTO_A_CONFERIR',
            'Confirmar com a Controladoria que o juizo/comarca aceita o ' + via,
            'Indeferido o diferimento, o prazo de emenda corre igual e a custa '
            'continua devida.',
            'Controladoria (RECOLHER CUSTAS)', bloqueia=False))

    # ---- via 3: custas recolhidas
    else:
        via = 'custas'
        tem_guia = bool(sit.get('guia'))
        tem_comprovante = bool(sit.get('comprovante'))
        conf = (conferir_documentos(nomes, DOCS_CUSTAS)
                if arquivos_pasta is not None and pasta_confiavel else None)
        if conf:
            achou = {d['slug'] for d in conf['ok']}
            tem_guia = tem_guia or 'guia-custas' in achou
            tem_comprovante = tem_comprovante or 'comprovante-custas' in achou

        if not tem_guia:
            pendencias.append(_pendencia(
                'SEM_GUIA',
                'Guia das custas iniciais nao foi emitida',
                'Sem topico de gratuidade a custa e devida na distribuicao. '
                'POP-CJ-006: a guia e emitida NO protocolo, nao depois dele.',
                'Controladoria',
                tarefa='RECOLHER CUSTAS'))
        if not tem_comprovante:
            pendencias.append(_pendencia(
                'SEM_COMPROVANTE',
                'Comprovante de pagamento das custas nao esta juntado',
                'Inicial sem custas e sem gratuidade gera emenda (art. 321 CPC) e, '
                'nao cumprida, indeferimento. Foi o que aconteceu no '
                '0000021-76.2026.8.22.0001.',
                'Controladoria + CS (cliente paga)',
                tarefa='RECOLHER CUSTAS'))
        if not sit.get('cs_acionado') and not tem_comprovante:
            pendencias.append(_pendencia(
                'CS_NAO_ACIONADO',
                'Cliente nao foi avisado do boleto das custas',
                'POP-CJ-006: a Controladoria emite e o CS leva ao cliente com '
                'valor, prazo e processo. Sem esse elo o boleto morre na pasta.',
                'CS (Anna Lydia / Karla)',
                tarefa='CONTATO COM CLIENTE'))
        if peca['anuncia_guia'] and not tem_comprovante:
            pendencias.append(_pendencia(
                'GUIA_ANUNCIADA_NAO_JUNTADA',
                'A peca lista a guia das custas no rol de documentos, mas nao ha '
                'guia paga nem juntada',
                'Rol que promete documento inexistente e o caminho curto para a '
                'emenda: o juizo le o rol, nao encontra o anexo e intima. '
                'Aconteceu no 0000021-76.2026.8.22.0001.',
                'Controladoria',
                tarefa='RECOLHER CUSTAS'))
        if not sit and arquivos_pasta is None:
            observacoes.append('Sem processo e sem pasta informados: a trava leu so o '
                               'texto da peca. Rode com --processo e --cliente.')

    if sit.get('emenda'):
        pendencias.append(_pendencia(
            'EMENDA_ABERTA',
            'Ja ha emenda a inicial registrada neste processo',
            'Emenda aberta quer dizer que o juizo ja apontou falta — conferir se '
            'o que falta e justamente a custa ou o documento desta conferencia.',
            'Controladoria', bloqueia=False))

    # ---- marcador de pendencia que sobrou na minuta
    for marcador in peca.get('marcadores_pendentes') or ():
        pendencias.append(_pendencia(
            'MARCADOR_PENDENTE',
            f'A peca ainda tem marcador de pendencia: "{marcador}"',
            'Marcador vivo quer dizer documento ou trecho que ficou faltando. '
            'A skill `timbrado` proibe "conforme documento anexo" sem documento.',
            'Advogado responsavel'))

    # ---- documentos da tese (mesmo rol da skill inicial-anexos)
    if rol_tese:
        for slug in rol_tese.get('faltando', ()):
            pendencias.append(_pendencia(
                'DOC_TESE_FALTANDO',
                f'Falta documento obrigatorio da tese: {slug}',
                'O rol da tese e o que instrui a inicial — faltando, a peca vai '
                'para emenda junto com as custas.',
                'Setor de documentos',
                tarefa='INICIAL - ORGANIZAR DOCUMENTOS'))

    # Pendencia que nasceu de olhar a pasta so bloqueia se a pasta for a certa.
    if not pasta_confiavel:
        for p in pendencias:
            if p['codigo'] in ('GRATUIDADE_SEM_DECLARACAO', 'GRATUIDADE_SEM_PROVA_RENDA',
                               'DOC_TESE_FALTANDO'):
                p['bloqueia'] = False
                p['o_que'] += ' (a conferir: a pasta lida pode nao ser a do cliente)'
        observacoes.append(
            'A pasta localizada no Drive nao parece ser a pasta de documentos do '
            'cliente. O que depende dela virou ressalva — confira a mao com: '
            'python OPERACIONAL/main.py drive cliente "NOME"')

    bloqueantes = [p for p in pendencias if p['bloqueia']]
    if motivo_excecao and bloqueantes:
        status = STATUS_LIBERADO
        observacoes.append('EXCECAO REGISTRADA: ' + motivo_excecao)
        observacoes.append('As pendencias abaixo continuam existindo — o protocolo foi '
                           'liberado por decisao humana, nao por elas terem sumido.')
    elif bloqueantes:
        status = STATUS_BLOQUEADO
    elif indefinida:
        status = STATUS_CONFERIR
    else:
        status = STATUS_LIBERADO

    return {'status': status, 'via': via, 'pendencias': pendencias,
            'observacoes': observacoes, 'excecao': motivo_excecao}


# ============================================================
# 7. PLANO DE TAREFAS (nada e' gravado aqui)
# ============================================================

def _tipos_tarefa():
    return (getattr(equipe, 'TIPOS_TAREFA', None) or {}) if equipe else {}


def _id_tarefa(nome):
    """ID do tipo de tarefa pelo nome, com as settings como fonte."""
    mapa = {'RECOLHER CUSTAS': 2947310,
            'CONTATO COM CLIENTE': 2940775,
            'INICIAL - ORGANIZAR DOCUMENTOS': 9970897,
            'SOLICITAR DOCUMENTACAO': 9970898}
    if advbox:
        try:
            achado = advbox.buscar_tipo_tarefa(nome)
            if achado:
                return achado
        except Exception:
            pass
    return mapa.get(nome)


def plano_tarefas(veredito, lawsuit):
    """
    Monta as tarefas que a pendencia pede — sem criar nenhuma.

    Quem lanca e' sempre a controller do advogado responsavel
    (roteamento_controller), como em toda tarefa da automacao.
    """
    if not lawsuit or veredito['status'] not in (STATUS_BLOQUEADO, STATUS_CONFERIR):
        return []

    rota = roteamento.resolver(lawsuit=lawsuit) if roteamento else {}
    from_id = rota.get('from_id')
    controller_id = rota.get('controller_id') or from_id
    comunicacao = ((getattr(equipe, 'COMUNICACAO_CLIENTE', None) or {}) if equipe else {})
    setor_docs = ((getattr(equipe, 'SETOR_PROVAS', None) or {}) if equipe else {})

    processo = lawsuit.get('process_number') or ''
    rodape = ('\n\n--\n[Trava de custas — POP-CJ-006 / POP-CJ-PROT-001] '
              'Conferencia automatica antes do protocolo. Confira antes de concluir.')

    tarefas, vistos = [], set()
    for p in veredito['pendencias']:
        if not p.get('tarefa') or p['tarefa'] in vistos:
            continue
        vistos.add(p['tarefa'])
        if p['tarefa'] == 'CONTATO COM CLIENTE':
            guests = [v for v in comunicacao.values() if v]
        elif p['tarefa'] in ('INICIAL - ORGANIZAR DOCUMENTOS', 'SOLICITAR DOCUMENTACAO'):
            guests = [v for v in setor_docs.values() if v]
        else:
            guests = [controller_id] if controller_id else []
        tarefas.append({
            'tipo': p['tarefa'],
            'tipo_id': _id_tarefa(p['tarefa']),
            'guests': [g for g in guests if g],
            'from_id': from_id,
            'lawsuit_id': lawsuit.get('id'),
            'comentario': (f"{p['o_que']} — processo {processo}.\n{p['por_que']}" + rodape),
        })
    return tarefas


def criar_tarefas(tarefas, confirmado=False):
    """Grava o plano no ADVBOX. Exige confirmacao explicita (regra de ouro)."""
    if not confirmado:
        raise ValueError('criar_tarefas exige confirmado=True')
    criadas = []
    for t in tarefas:
        if not (t['tipo_id'] and t['guests'] and t['from_id'] and t['lawsuit_id']):
            criadas.append({'tarefa': t['tipo'], 'ok': False,
                            'erro': 'faltou tipo, destinatario, remetente ou processo'})
            continue
        try:
            r = advbox.criar_publicacao(
                lawsuit_id=t['lawsuit_id'], task_id=t['tipo_id'], guest_ids=t['guests'],
                comments=t['comentario'], from_id=str(t['from_id']), urgent=True)
            criadas.append({'tarefa': t['tipo'], 'ok': bool(r), 'resposta': r})
        except Exception as e:
            criadas.append({'tarefa': t['tipo'], 'ok': False, 'erro': str(e)})
    return criadas


# ============================================================
# 8. CAMADA 2 — AUDITORIA DO QUE JA FOI PROTOCOLADO
# ============================================================

_STAGE_PROTOCOLADA = ('acao protocolada', 'postulatoria', 'distribuid')

# NNNNNNN-DD.AAAA.J.TR.OOOO — o ultimo grupo e' a unidade de origem, e '0000'
# quer dizer que o processo nasceu no proprio tribunal (agravo, mandado de
# seguranca originario). Recurso tem PREPARO, que e' outra regua e outro POP:
# deixar agravo na trava de custas iniciais so produz alarme falso, e alarme
# falso e o que faz a controller parar de ler o relatorio.
_CNJ = re.compile(r'(\d{7})-(\d{2})\.(\d{4})\.(\d)\.(\d{2})\.(\d{4})')


def analisar_numero(numero):
    m = _CNJ.search(numero or '')
    if not m:
        return {'valido': False, 'ano': None, 'segundo_grau': False}
    return {'valido': True, 'ano': int(m.group(3)), 'segmento': m.group(4),
            'tribunal': m.group(5), 'origem': m.group(6),
            'segundo_grau': m.group(6) == '0000'}


def escopo_da_trava(lawsuit, hoje=None):
    """Diz se o processo e' uma inicial de 1o grau recente (o que a trava mede).

    Devolve (dentro_do_escopo, motivo)."""
    hoje = hoje or date.today()
    numero = lawsuit.get('process_number') or ''
    info = analisar_numero(numero)
    if not numero.strip():
        return True, 'ainda sem numero — inicial aguardando distribuicao'
    if not info['valido']:
        return False, 'numero fora do padrao CNJ — conferir o cadastro'
    if info['segundo_grau']:
        return False, 'processo originario do tribunal (recurso) — preparo, nao custa inicial'
    if info['ano'] < hoje.year:
        return False, (f"processo de {info['ano']} cadastrado agora — "
                       'cadastro tardio, nao distribuicao nova')
    return True, 'inicial de 1o grau do ano corrente'


def _data(valor):
    if not valor:
        return None
    for formato in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(valor)[:19], formato).date()
        except ValueError:
            continue
    return None


def protocolos_concluidos(dias=7, hoje=None):
    """IDs de processo cujas tarefas de PROTOCOLO foram CONCLUIDAS no periodo.

    E' por aqui que a auditoria acha o protocolo de verdade. Procurar so por
    processo cadastrado nos ultimos dias nao serve: o caso entra no ADVBOX quando
    o contrato e' assinado e so vai a protocolo meses depois. O alongamento do
    Sr. Cliente Y foi cadastrado em 21/07/2026, protocolado em 22/09 e
    escapou das duas camadas — da fila, porque a tarefa ja estava concluida; da
    auditoria, porque o cadastro nao era recente.
    """
    hoje = hoje or date.today()
    inicio = hoje - timedelta(days=int(dias))
    achados = {}
    for rotulo, tid in TAREFAS_FILA.items():
        try:
            tarefas = advbox.listar_tarefas(task_id=tid, completed_start=inicio.isoformat(),
                                            completed_end=hoje.isoformat(), limit=100) or []
        except Exception:
            continue
        for t in tarefas:
            lid = t.get('lawsuits_id') or (t.get('lawsuit') or {}).get('id')
            if not lid:
                continue
            quando = max((u.get('completed') or '') for u in (t.get('users') or [{}]))
            anterior = achados.get(lid)
            if not anterior or (quando and quando > anterior[1]):
                achados[lid] = (rotulo, quando or '', t.get('notes') or '')
    return achados


def iniciais_recentes(dias=7, hoje=None, processos=None):
    """Processos a auditar: os que tiveram PROTOCOLO concluido no periodo e,
    como rede extra, os cadastrados ha pouco e ja em fase protocolada."""
    hoje = hoje or date.today()
    if processos is None:
        processos = advbox.listar_processos() or []
    protocolados = protocolos_concluidos(dias=dias, hoje=hoje)
    recentes, fora = [], []
    for p in processos:
        criado = _data(p.get('created_at'))
        marca = protocolados.get(p.get('id'))
        if not marca and (not criado or (hoje - criado).days > dias):
            continue
        if marca:
            p = dict(p, _protocolo_tarefa=marca[0], _protocolo_em=marca[1],
                     _protocolo_nota=marca[2])
        estagio = _norm(p.get('stage') or '')
        if not marca and estagio and not any(m in estagio for m in _STAGE_PROTOCOLADA):
            continue
        dentro, motivo = escopo_da_trava(p, hoje=hoje)
        (recentes if dentro else fora).append(
            dict(p, _motivo_escopo=motivo) if not dentro else p)
    recentes.sort(key=lambda p: p.get('_protocolo_em') or p.get('created_at') or '',
                  reverse=True)
    return recentes, fora


def auditar_processo(lawsuit, hoje=None):
    """Confere um processo ja' protocolado: ha' sinal de custas ou de gratuidade?"""
    hoje = hoje or date.today()
    sit = situacao_advbox(lawsuit.get('id'))
    criado = _data(lawsuit.get('created_at'))
    dias = (hoje - criado).days if criado else None

    tem_comprovante = bool(sit.get('comprovante'))
    tem_guia = bool(sit.get('guia'))
    emenda = bool(sit.get('emenda'))

    if tem_comprovante:
        status, motivo = 'RESOLVIDO', 'comprovante de pagamento registrado no ADVBOX'
    elif emenda:
        status, motivo = ('PROTOCOLADO SEM CUSTAS',
                          'ja ha emenda a inicial aberta — a custa foi cobrada pelo juizo')
    elif tem_guia:
        status, motivo = ('CUSTAS EM ABERTO',
                          'guia emitida, pagamento ainda nao registrado')
    else:
        status, motivo = ('SEM SINAL DE CUSTAS',
                          'nenhum registro de guia, pagamento ou gratuidade')

    return {'lawsuit': lawsuit, 'situacao': sit, 'status': status, 'motivo': motivo,
            'dias_desde_cadastro': dias,
            'guia': sit.get('guia'), 'comprovante': sit.get('comprovante'),
            'emenda': sit.get('emenda')}


def auditar(dias=7, hoje=None, limite=60, processos=None):
    """Rede de seguranca: varre as iniciais recentes e devolve as sem custas."""
    candidatos, fora = iniciais_recentes(dias=dias, hoje=hoje, processos=processos)
    estourou = len(candidatos) > limite
    resultado = [auditar_processo(p, hoje=hoje) for p in candidatos[:limite]]
    return {'total': len(candidatos), 'analisados': len(resultado),
            'estourou_limite': estourou, 'itens': resultado, 'fora_escopo': fora}


# ============================================================
# 8b. CAMADA 0 — A FILA DE PROTOCOLO (o que roda sozinho)
# ============================================================
# A trava preventiva so fecha o ciclo se alguem a rodar. Fechar pelo arquivo da
# peca nao da: conferido em 22/09/2026, as tarefas de protocolo do ADVBOX NAO
# trazem a peca anexada (`GET /documents?post_id=` volta vazio nas 66 tarefas
# PROTOCOLO D-3 abertas). O que existe, e e' suficiente para a conferencia, e' a
# propria fila: tipo de tarefa, rotulo do que sera protocolado e o historico de
# custas do processo.
#
# Por isso a rodada das 08:00 le a FILA, nao a peca: para cada item que vai a
# protocolo nos proximos dias, pergunta se ha guia/comprovante/gratuidade
# registrados. O texto da peca so entra quando a controller roda
# `custas conferir` no arquivo — ai a leitura e' completa.

TAREFAS_FILA = {
    'PECA APROVADA PARA PROTOCOLO': 8720914,
    'PROTOCOLO D-3': 8941168,
    'PROTOCOLO D-2': 10177444,
    'PROTOCOLO D-1': 10177445,
    'PROTOCOLO - PRAZO FATAL': 10177446,
}

# Rotulos como a controller escreve na tarefa ("Contrarrazoes", "inicial
# anulatoria", "descaracterizacao da mora"). Ordem importa: o primeiro que casar
# manda, e os que NAO pagam vem antes para "contrarrazoes de apelacao" nao virar
# apelacao.
_ROTULOS = (
    ('sem-preparo', r'contrarrazoes|contraminuta|embargos de declaracao|agravo interno|'
                    r'agravo regimental|memoriais|sustentacao'),
    ('recurso', r'agravo de instrumento|\bapelacao\b|recurso inominado|'
                r'recurso especial|recurso extraordinario|\bagravo\b'),
    ('defesa-execucao', r'embargos a execucao|excecao de pre-?executividade|'
                        r'impugnacao ao cumprimento'),
    # "inicial" no rotulo decide antes de qualquer outra coisa.
    ('inicial', r'\binicial\b'),
    # Peca intermediaria: nao paga custa de distribuicao nem preparo.
    ('outra', r'manifestacao|replica|impugnacao|alegacoes finais|habilitacao|'
              r'substabelecimento|memorial|oficio|cumprimento de sentenca|'
              r'especificacao de provas|quesitos|\bpeticao\b'),
    # Nome da acao. `\bacao\b` com fronteira: sem ela, "documentACAO pendente"
    # fazia "Protocolar manifestacao. (documentacao pendente)" virar inicial —
    # apareceu na fila real de 22/09/2026.
    ('inicial', r'\bacao\b|\bacoes\b|declaratoria|anulatoria|revisional|alongamento|'
                r'descaracterizacao|monitoria|obrigacao de fazer|indenizatoria|'
                r'mandado de seguranca|busca e apreensao'),
)


def classificar_rotulo(texto):
    """Natureza do que sera protocolado, pelo rotulo da tarefa."""
    t = _norm(texto)
    if not t.strip():
        return 'indefinido'
    for natureza, padrao in _ROTULOS:
        if re.search(padrao, t):
            return natureza
    return 'indefinido'


def _tarefas_da_fila(dias, hoje=None):
    hoje = hoje or date.today()
    inicio = (hoje - timedelta(days=2)).isoformat()
    fim = (hoje + timedelta(days=dias)).isoformat()
    itens, vistos = [], set()
    for nome, tid in TAREFAS_FILA.items():
        try:
            achados = advbox.listar_tarefas(task_id=tid, date_start=inicio,
                                            date_end=fim) or []
        except Exception as e:
            print(f'  AVISO: nao consegui ler a fila "{nome}": {e}')
            continue
        for t in achados:
            if t.get('id') in vistos:
                continue
            vistos.add(t['id'])
            itens.append(dict(t, _fila=nome))
    itens.sort(key=lambda t: t.get('date') or '')
    return itens


def _indice_processos(processos=None):
    """{lawsuit_id: cadastro completo} — o processo que vem DENTRO da tarefa
    (`/posts`) traz so id, numero e cliente: **nao traz `responsible_id`**, e sem
    ele `roteamento_controller` cai no fallback e manda tudo para a mesma
    controller (aconteceu na 1a rodada da fila, 22/09/2026: 37 tarefas para a
    Manuelle). A carteira inteira custa 3 GET (limit=1000) e resolve todos.
    """
    if processos is None:
        processos = advbox.listar_processos() or []
    return {p['id']: p for p in processos if p.get('id')}


def conferir_fila(dias=5, hoje=None, tarefas=None, processos=None):
    """
    Passa a fila de protocolo dos proximos dias pela trava.

    Somente leitura. Item que nao e' inicial nem recurso com preparo nao gera
    consulta ao /history (o teto e 30 GET/min).
    """
    hoje = hoje or date.today()
    itens = tarefas if tarefas is not None else _tarefas_da_fila(dias, hoje=hoje)
    indice = None
    saida = []
    for t in itens:
        law = t.get('lawsuit') or {}
        rotulo = t.get('notes') or ''
        natureza = classificar_rotulo(rotulo)
        # O cadastro completo (com responsavel) vale para todo item que a
        # controller vai ler, nao so para os que consultam o /history.
        if natureza != 'sem-preparo' and law.get('id'):
            if indice is None:
                indice = _indice_processos(processos)
            law = indice.get(law['id']) or law

        registro = {'tarefa': t, 'fila': t.get('_fila'), 'rotulo': rotulo,
                    'natureza': natureza, 'lawsuit': law,
                    'processo': law.get('process_number'),
                    'situacao': None, 'status': None, 'motivo': ''}

        if natureza == 'sem-preparo':
            registro['status'] = 'FORA DA TRAVA'
            registro['motivo'] = ('nao paga custa inicial nem preparo '
                                  '(ED: art. 1.023; contrarrazoes nao e recurso)')
        elif natureza == 'outra':
            # Peca intermediaria (manifestacao, replica, alegacoes finais) nao paga
            # custa de distribuicao nem preparo. Na 1a rodada ela caia no ramo das
            # custas e saia como "SEM REGISTRO - custas iniciais": 14 dos 32
            # alarmes da fila eram isso.
            registro['status'] = 'FORA DA TRAVA'
            registro['motivo'] = 'peca intermediaria: sem custa de distribuicao nem preparo'
        elif natureza == 'defesa-execucao':
            registro['status'] = 'A CONFERIR'
            registro['motivo'] = 'custas conforme o regimento do tribunal'
        elif natureza == 'indefinido':
            registro['status'] = 'A CONFERIR'
            registro['motivo'] = ('nao da para dizer pelo rotulo o que sera protocolado'
                                  if rotulo.strip() else 'tarefa sem descricao')
        else:
            if not law.get('id'):
                registro['status'] = 'SEM CADASTRO'
                registro['motivo'] = 'tarefa sem processo vinculado no ADVBOX'
            else:
                sit = situacao_advbox(law['id'])
                registro['situacao'] = sit
                exigencia = 'preparo' if natureza == 'recurso' else 'custas iniciais'
                if sit.get('comprovante'):
                    registro['status'] = 'OK'
                    registro['motivo'] = f'{exigencia}: comprovante registrado'
                elif sit.get('guia'):
                    registro['status'] = 'EM ABERTO'
                    registro['motivo'] = f'{exigencia}: guia emitida, pagamento nao registrado'
                else:
                    registro['status'] = 'SEM REGISTRO'
                    registro['motivo'] = (f'{exigencia}: nenhuma guia, pagamento ou '
                                          'gratuidade registrada')
        saida.append(registro)
    return saida


def pendencia_da_fila(registro):
    """A pendencia que o item da fila merece — para virar tarefa, se a controller quiser."""
    if registro['status'] not in ('SEM REGISTRO', 'EM ABERTO'):
        return None
    if registro['natureza'] == 'recurso':
        return _pendencia(
            'FILA_PREPARO',
            f"Recurso na fila de protocolo sem preparo comprovado ({registro['rotulo'][:60]})",
            'art. 1.007, caput e § 4o: sem comprovacao no ato da interposicao o '
            'recolhimento passa a ser EM DOBRO, vedada a complementacao (§ 5o). '
            'Recurso nao tem emenda.',
            'Controladoria', tarefa='RECOLHER CUSTAS')
    return _pendencia(
        'FILA_CUSTAS',
        f"Inicial na fila de protocolo sem custas registradas ({registro['rotulo'][:60]})",
        'art. 321: protocolada sem custas e sem gratuidade, o juizo intima para '
        'emendar em 15 dias e, nao cumprido, indefere a inicial.',
        'Controladoria', tarefa='RECOLHER CUSTAS')


# ============================================================
# 9. AUTOTESTE
# ============================================================

_CASOS = [
    # (texto, esperado_tipo, esperado_gratuidade)
    ('EXCELENTISSIMO SENHOR DOUTOR JUIZ DE DIREITO DA 1a VARA CIVEL DA COMARCA DE '
     'PORTO VELHO/RO. CLIENTE G ALVES, brasileira, produtora rural, inscrita no CPF, '
     'residente e domiciliada na Linha 05, vem propor a presente ACAO DECLARATORIA. '
     'Requer a citacao da parte requerida. Da-se a causa o valor de R$ 1.000,00.',
     'inicial', False),
    ('EXCELENTISSIMO JUIZ. JOAO, ja qualificado nos autos em epigrafe, vem apresentar '
     'REPLICA a contestacao apresentada pelo banco.', 'outra', False),
    ('EGREGIO TRIBUNAL DE JUSTICA. Razoes de apelacao. O apelante requer a reforma.',
     'outra', False),
    ('Vem ajuizar a presente acao de descaracterizacao da mora. Requer a citacao do '
     'requerido. Requer, ainda, os beneficios da gratuidade da justica, por ser '
     'hipossuficiente. Da-se a causa o valor de R$ 1.000,00.', 'inicial', True),
    ('Vem propor a presente acao. Requer a concessao do beneficio previsto no art. 98 '
     'do CPC, isentando-se do recolhimento das custas. Requer a citacao.', 'inicial', True),
    ('Trata-se de manifestacao sobre os documentos juntados, nos autos em epigrafe.',
     'outra', False),
    ('Peticao inicial. Vem propor acao. Requer a citacao. Da-se a causa o valor de '
     'R$ 50.000,00. Requer o diferimento das custas ao final.', 'inicial', False),
]


def _autoteste():
    falhas = 0
    for texto, tipo_esperado, grat_esperada in _CASOS:
        r = analisar_peca(texto)
        if r['tipo'] != tipo_esperado:
            print(f"  [X] tipo: esperado {tipo_esperado}, veio {r['tipo']} — {texto[:60]}...")
            falhas += 1
        if r['gratuidade'] != grat_esperada:
            print(f"  [X] gratuidade: esperado {grat_esperada}, veio {r['gratuidade']}"
                  f" — {texto[:60]}...")
            falhas += 1

    # diferimento reconhecido
    r = analisar_peca(_CASOS[6][0])
    if not r['diferimento']:
        print('  [X] nao reconheceu o pedido de diferimento'); falhas += 1
    if r['valor_causa'] != '50.000,00':
        print(f"  [X] valor da causa: veio {r['valor_causa']}"); falhas += 1

    # veredito: inicial sem gratuidade e sem custas -> BLOQUEADO
    v = avaliar(peca=analisar_peca(_CASOS[0][0]), arquivos_pasta=[], advbox_situacao={})
    if v['status'] != STATUS_BLOQUEADO:
        print(f"  [X] inicial sem custas devia bloquear, veio {v['status']}"); falhas += 1
    if not any(p['codigo'] == 'SEM_COMPROVANTE' for p in v['pendencias']):
        print('  [X] faltou a pendencia SEM_COMPROVANTE'); falhas += 1

    # com comprovante juntado -> LIBERADO
    v = avaliar(peca=analisar_peca(_CASOS[0][0]),
                arquivos_pasta=['Guia custas iniciais.pdf', 'comprovante de pagamento.pdf'],
                advbox_situacao={'guia': {'task': 'RECOLHER CUSTAS'},
                                 'comprovante': {'task': 'COMENTARIO',
                                                 'comments': 'comprovante de pagamento'}})
    if v['status'] != STATUS_LIBERADO:
        print(f"  [X] inicial com custas pagas devia liberar, veio {v['status']}"
              f" ({[p['codigo'] for p in v['pendencias']]})"); falhas += 1

    # gratuidade sem declaracao na pasta -> BLOQUEADO
    v = avaliar(peca=analisar_peca(_CASOS[3][0]), arquivos_pasta=['procuracao.pdf'],
                advbox_situacao={})
    if v['status'] != STATUS_BLOQUEADO:
        print(f"  [X] gratuidade sem declaracao devia bloquear, veio {v['status']}"); falhas += 1

    # gratuidade com declaracao -> LIBERADO (prova de renda e ressalva, nao bloqueio)
    v = avaliar(peca=analisar_peca(_CASOS[3][0]),
                arquivos_pasta=['DECLARACAO DE HIPOSSUFICIENCIA.pdf'], advbox_situacao={})
    if v['status'] != STATUS_LIBERADO:
        print(f"  [X] gratuidade com declaracao devia liberar, veio {v['status']}"); falhas += 1

    # peca que nao e inicial -> trava nao se aplica
    v = avaliar(peca=analisar_peca(_CASOS[1][0]), arquivos_pasta=[], advbox_situacao={})
    if v['status'] != STATUS_NAO_SE_APLICA:
        print(f"  [X] replica nao devia entrar na trava, veio {v['status']}"); falhas += 1

    # excecao registrada libera, mas mantem as pendencias na tela
    v = avaliar(peca=analisar_peca(_CASOS[0][0]), arquivos_pasta=[], advbox_situacao={},
                motivo_excecao='prazo fatal hoje — Dra. Juliana autorizou')
    if v['status'] != STATUS_LIBERADO or not v['pendencias']:
        print('  [X] excecao devia liberar mantendo as pendencias'); falhas += 1

    # caso real: a peca ANUNCIA a guia no rol e ela nao existe -> bloqueia
    texto_rosiane = (
        'EXCELENTISSIMO JUIZ DE DIREITO DA VARA CIVEL DE PORTO VELHO. CLIENTE G, '
        'produtora rural, qualificada, residente na Linha 05, vem propor a presente '
        'acao declaratoria de descaracterizacao da mora. Requer a citacao do requerido. '
        'Instruem a inicial a procuracao, a cedula, o demonstrativo de conta vinculada '
        'e guia de recolhimento das custas iniciais. Da-se a causa o valor de '
        'R$ 1.000,00 (mil reais).')
    r = analisar_peca(texto_rosiane)
    if not r['anuncia_guia'] or r['gratuidade'] or r['tipo'] != 'inicial':
        print(f'  [X] leitura do caso Cliente G: {r["tipo"]}, anuncia_guia='
              f'{r["anuncia_guia"]}, gratuidade={r["gratuidade"]}'); falhas += 1
    v = avaliar(peca=r, arquivos_pasta=['procuracao.pdf', 'cedula.pdf'], advbox_situacao={})
    if v['status'] != STATUS_BLOQUEADO:
        print(f'  [X] o caso Cliente G devia bloquear, veio {v["status"]}'); falhas += 1
    if not any(p['codigo'] == 'GUIA_ANUNCIADA_NAO_JUNTADA' for p in v['pendencias']):
        print('  [X] faltou a pendencia GUIA_ANUNCIADA_NAO_JUNTADA'); falhas += 1

    # marcador de pendencia vivo na minuta bloqueia, mesmo com as custas pagas
    r = analisar_peca(_CASOS[0][0] + ' [PENDENTE] declaracao de hipossuficiencia assinada.')
    v = avaliar(peca=r,
                arquivos_pasta=['guia.pdf', 'comprovante de pagamento.pdf'],
                advbox_situacao={'guia': {'task': 'RECOLHER CUSTAS'},
                                 'comprovante': {'task': 'COMENTARIO',
                                                 'comments': 'comprovante de pagamento'}})
    if v['status'] != STATUS_BLOQUEADO or not any(
            p['codigo'] == 'MARCADOR_PENDENTE' for p in v['pendencias']):
        print(f'  [X] marcador [PENDENTE] devia bloquear, veio {v["status"]}'); falhas += 1

    # rotulos da fila real do ADVBOX (tarefas PROTOCOLO D-3 abertas em 22/09/2026)
    for rotulo, esperado in (
            ('Contrarrazoes', 'sem-preparo'),
            ('Embargos À Execução', 'defesa-execucao'),
            ('Manifestação', 'outra'),
            ('Réplica à contestação', 'outra'),
            ('Alegações finais', 'outra'),
            ('Protocolar manifestação. (documentação pendente).', 'outra'),
            ('Protocolar ação de descaracterização da mora.', 'inicial'),
            ('Protocolar recurso especial.', 'recurso'),
            ('inicial de descaracterização da mora - documentos pendentes', 'inicial'),
            ('inicial anulatória', 'inicial'),
            ('descaracterização da mora', 'inicial'),
            ('descaracterização da mora da CEF', 'inicial'),
            ('Agravo de Instrumento', 'recurso'),
            ('Apelação', 'recurso'),
            ('Contrarrazões de apelação', 'sem-preparo'),
            ('Embargos de Declaração', 'sem-preparo'),
            ('', 'indefinido')):
        veio = classificar_rotulo(rotulo)
        if veio != esperado:
            print(f'  [X] rotulo "{rotulo}": esperado {esperado}, veio {veio}'); falhas += 1

    # recurso: preparo, desercao e as pecas que NAO pagam preparo
    agravo = ('EXCELENTISSIMO DESEMBARGADOR RELATOR. FULANO, ja qualificado nos autos, '
              'vem, respeitosamente, com fundamento no art. 1.015, I, do CPC, interpor '
              'o presente AGRAVO DE INSTRUMENTO contra a decisao que indeferiu a tutela.')
    r = analisar_peca(agravo)
    if r['tipo'] != 'recurso' or r['recurso'] != 'agravo-instrumento' or not r['preparo']:
        print(f"  [X] agravo: veio {r['tipo']}/{r['recurso']}"); falhas += 1
    v = avaliar(peca=r, advbox_situacao={})
    if v['status'] != STATUS_BLOQUEADO or v['via'] != 'preparo':
        print(f"  [X] agravo sem preparo devia bloquear, veio {v['status']}"); falhas += 1
    if not any(p['codigo'] == 'AGRAVO_SEM_COMPROVANTE_NA_PECA' for p in v['pendencias']):
        print('  [X] faltou a exigencia do art. 1.017, § 1o'); falhas += 1

    apelacao = ('EXCELENTISSIMO JUIZ. FULANO, ja qualificado nos autos em epigrafe, vem '
                'interpor a presente APELACAO contra a sentenca de improcedencia.')
    r = analisar_peca(apelacao)
    if r['tipo'] != 'recurso' or r['recurso'] != 'apelacao':
        print(f"  [X] apelacao: veio {r['tipo']}/{r['recurso']}"); falhas += 1
    v = avaliar(peca=r, advbox_situacao={'guia': {'task': 'RECOLHER CUSTAS'},
                                         'comprovante': {'comments': 'comprovante de pagamento'}})
    if v['status'] != STATUS_LIBERADO:
        print(f"  [X] apelacao com preparo pago devia liberar, veio {v['status']}"); falhas += 1

    for texto, esperado in (
            ('FULANO vem opor os presentes EMBARGOS DE DECLARACAO contra a sentenca.',
             'embargos-declaracao'),
            ('FULANO vem apresentar CONTRARRAZOES ao recurso de apelacao interposto.',
             'contrarrazoes'),
            ('FULANO vem interpor AGRAVO INTERNO contra a decisao monocratica.',
             'agravo-interno')):
        r = analisar_peca(texto)
        if r['recurso'] != esperado:
            print(f'  [X] {esperado}: veio {r["recurso"]}'); falhas += 1
        v = avaliar(peca=r, advbox_situacao={})
        if v['status'] != STATUS_NAO_SE_APLICA:
            print(f'  [X] {esperado} nao paga preparo, devia ficar fora: {v["status"]}')
            falhas += 1

    # nome de recurso citado no corpo nao transforma a peca em recurso
    r = analisar_peca('Vem propor a presente acao declaratoria. Requer a citacao. '
                      'Da-se a causa o valor de R$ 1.000,00. A Orientacao 4 do REsp '
                      '1.061.530/RS e a apelacao citada nao se aplicam aqui.')
    if r['tipo'] != 'inicial' or r['recurso']:
        print(f"  [X] inicial que cita REsp virou {r['tipo']}/{r['recurso']}"); falhas += 1

    # defesa em execucao: nem inicial, nem preparo
    r = analisar_peca('FULANO, ja qualificado, opoe os presentes EMBARGOS A EXECUCAO '
                      'em face do banco, nos autos da execucao em epigrafe.')
    v = avaliar(peca=r, advbox_situacao={})
    if v['status'] != STATUS_CONFERIR:
        print(f"  [X] embargos a execucao deviam ir para conferencia: {v['status']}")
        falhas += 1

    # pasta errada: o que depende dela vira ressalva, o que vem do ADVBOX continua bloqueando
    for nome_pasta, esperado in (
            ('FULANA DE TAL', True),
            ('FULANA DE TAL - SEM PROCESSO', False),
            ('FULANA DE TAL - 0000021-76.2026.8.22.0001', False),
            ('FULANA DE TAL ', True)):
        if pasta_confere(nome_pasta, 'FULANA DE TAL') != esperado:
            print(f'  [X] pasta_confere("{nome_pasta}") devia ser {esperado}'); falhas += 1
    v = avaliar(peca=analisar_peca(_CASOS[3][0]), arquivos_pasta=['peca.docx'],
                advbox_situacao={}, pasta_confiavel=False)
    if any(p['codigo'] == 'GRATUIDADE_SEM_DECLARACAO' and p['bloqueia']
           for p in v['pendencias']):
        print('  [X] pasta duvidosa nao devia bloquear por documento'); falhas += 1

    # escopo: agravo e processo antigo ficam fora da trava
    hoje_teste = date(2026, 9, 22)
    casos_escopo = [
        ({'process_number': '0000021-76.2026.8.22.0001'}, True),
        ({'process_number': '0000039-03.2026.8.22.0000'}, False),   # agravo (origem 0000)
        ({'process_number': '0000040-69.1995.8.22.0001'}, False),   # de 1995
        ({'process_number': ''}, True),                             # sem numero ainda
    ]
    for law, esperado in casos_escopo:
        dentro, motivo = escopo_da_trava(law, hoje=hoje_teste)
        if dentro != esperado:
            print(f"  [X] escopo de {law['process_number'] or '(sem numero)'}: "
                  f'esperado {esperado}, veio {dentro} ({motivo})'); falhas += 1

    # eventos do ADVBOX do caso real
    hist = [
        {'task': 'RECOLHER CUSTAS', 'comments': 'Emitir custas para pagamento.'},
        {'task': 'CONTATO COM CLIENTE',
         'comments': 'Entrar em contato com a cliente enviar boleto referente as custas.'},
        {'task': 'COMENTARIO', 'comments': 'Segue comprovante de pagamento:'},
        {'task': 'SOLICITAR DOCUMENTACAO',
         'comments': 'Solicitar documentos para emenda (comprovante de residencia e '
                     'juntada de custas).'},
    ]
    sit = situacao_advbox(None, historico=hist, pendentes=[])
    for chave in ('guia', 'cs_acionado', 'comprovante', 'emenda'):
        if not sit.get(chave):
            print(f'  [X] nao reconheceu o evento {chave} do caso Cliente G'); falhas += 1

    print(f"\n  {len(_CASOS)} casos de texto + vereditos: "
          f"{'TODOS OK' if not falhas else str(falhas) + ' FALHA(S)'}")
    return falhas


if __name__ == '__main__':
    if '--autoteste' in sys.argv:
        sys.exit(1 if _autoteste() else 0)
    print(__doc__)
